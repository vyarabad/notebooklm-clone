"""Text extraction for each supported source type.

Every extractor returns a list of Sections. A Section is a piece of text plus
a location label (e.g. "page 3", "slide 2") that is kept as citation metadata.
"""
import re
from dataclasses import dataclass
from pathlib import Path
from urllib.parse import urlparse

import requests
from bs4 import BeautifulSoup

SUPPORTED_EXTENSIONS = {".pdf", ".pptx", ".txt"}


class IngestionError(Exception):
    """Raised with a user-friendly message when a source can't be read."""


@dataclass
class Section:
    text: str
    location: str = ""  # "page 3", "slide 2", or "" for TXT/URL


def _clean(text: str) -> str:
    text = text.replace("\x00", " ")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n\s*\n\s*\n+", "\n\n", text)
    return text.strip()


def extract_pdf(path: str) -> list[Section]:
    from pypdf import PdfReader
    try:
        reader = PdfReader(path)
    except Exception as e:
        raise IngestionError(f"Couldn't open PDF: {e}")
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception:
            raise IngestionError("This PDF is password protected.")
    sections = []
    for i, page in enumerate(reader.pages, start=1):
        try:
            text = _clean(page.extract_text() or "")
        except Exception:
            text = ""
        if text:
            sections.append(Section(text, f"page {i}"))
    if not sections:
        raise IngestionError("No text found in this PDF. It may be a scanned image (OCR is not supported).")
    return sections


def extract_pptx(path: str) -> list[Section]:
    from pptx import Presentation
    try:
        prs = Presentation(path)
    except Exception as e:
        raise IngestionError(f"Couldn't open PPTX: {e}")
    sections = []
    for i, slide in enumerate(prs.slides, start=1):
        parts = []
        for shape in slide.shapes:
            if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                parts.append(shape.text_frame.text)
            if getattr(shape, "has_table", False) and shape.has_table:
                for row in shape.table.rows:
                    parts.append(" | ".join(cell.text for cell in row.cells))
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text if slide.notes_slide.notes_text_frame else ""
            if notes.strip():
                parts.append(f"Speaker notes: {notes}")
        text = _clean("\n".join(parts))
        if text:
            sections.append(Section(text, f"slide {i}"))
    if not sections:
        raise IngestionError("No text found in this PPTX.")
    return sections


def extract_txt(path: str) -> list[Section]:
    raw = Path(path).read_bytes()
    for enc in ("utf-8", "utf-16", "latin-1"):
        try:
            text = raw.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    text = _clean(text)
    if not text:
        raise IngestionError("This TXT file is empty.")
    return [Section(text, "")]


def extract_file(path: str) -> list[Section]:
    ext = Path(path).suffix.lower()
    if ext == ".pdf":
        return extract_pdf(path)
    if ext == ".pptx":
        return extract_pptx(path)
    if ext == ".txt":
        return extract_txt(path)
    raise IngestionError(f"Unsupported file type '{ext}'. Use PDF, PPTX, or TXT.")


def extract_url(url: str) -> tuple[str, list[Section]]:
    """Fetch a single web page. Returns (page_title, sections)."""
    url = (url or "").strip()
    parsed = urlparse(url)
    if parsed.scheme not in ("http", "https") or not parsed.netloc:
        raise IngestionError("Please enter a full URL starting with http:// or https://")
    try:
        resp = requests.get(url, timeout=20, headers={
            "User-Agent": "Mozilla/5.0 (compatible; NotebookLMClone/1.0)"
        })
        resp.raise_for_status()
    except requests.RequestException as e:
        raise IngestionError(f"Couldn't fetch the URL: {e}")

    ctype = resp.headers.get("Content-Type", "")
    if "html" not in ctype and "text" not in ctype:
        raise IngestionError(f"URL didn't return a web page (content type: {ctype or 'unknown'}).")

    soup = BeautifulSoup(resp.text, "html.parser")
    title = (soup.title.string.strip() if soup.title and soup.title.string else parsed.netloc)
    for tag in soup(["script", "style", "noscript", "nav", "footer", "header", "aside", "form", "svg"]):
        tag.decompose()
    main = soup.find("article") or soup.find("main") or soup.body or soup
    text = _clean(main.get_text("\n"))
    # drop very short lines (menus, buttons)
    lines = [ln for ln in text.split("\n") if len(ln.strip()) > 2]
    text = _clean("\n".join(lines))
    if len(text) < 100:
        raise IngestionError("Couldn't find readable text on that page (it may need JavaScript or a login).")
    return title[:150], [Section(text, "")]
